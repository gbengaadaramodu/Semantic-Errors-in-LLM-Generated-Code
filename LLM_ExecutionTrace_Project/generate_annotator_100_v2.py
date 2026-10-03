import json
import openpyxl
from openpyxl.styles import Alignment, Font
from pathlib import Path
from objective2.config import Objective2Config
from objective2.pipeline import Objective2Pipeline

# Load the exact 100 cases
with open('annotation/Validation_Corpus_100.json', 'r', encoding='utf-8') as f:
    corpus_100 = json.load(f)
target_trace_ids = {r['trace_id'] for r in corpus_100}

config = Objective2Config()
pipe = Objective2Pipeline(config)
all_records = pipe.corpus_loader.load_records()
eligible_records = [r for r in all_records if r.trace_id in target_trace_ids]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Annotation_Task_100"

headers = [
    "Trace ID", 
    "Dataset", 
    "Problem Summary (Prompt)", 
    "Executed Script (Generated Code + Test Suite)",
    "Traceback / Error", 
    "-> YOUR ANNOTATION: Fault Line (Number)", 
    "-> YOUR ANNOTATION: Causal Variable", 
    "-> YOUR ANNOTATION: Bug Description"
]
ws.append(headers)

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = Font(bold=True)
    if "YOUR ANNOTATION" in header:
        cell.font = Font(bold=True, color="0000FF")

def add_line_numbers(code: str) -> str:
    if not code:
        return ""
    lines = code.splitlines()
    return "\n".join([f"{i+1:02d} | {line}" for i, line in enumerate(lines)])

for r in eligible_records:
    full_record = pipe.trace_loader.enrich_record(r)
    
    parts = []
    dataset = full_record.dataset.strip().upper()
    
    # We must assemble the script EXACTLY as it was executed to produce the traceback.
    # From executor.py logic:
    if dataset == "MBPP":
        if full_record.test_imports:
            parts.append(full_record.test_imports)
            parts.append("\n\n")
        parts.append(full_record.generated_code)
        parts.append("\n\n")
        parts.append(full_record.test_code)
        parts.append("\n")
    else: # HumanEval
        parts.append(full_record.generated_code)
        parts.append("\n\n")
        parts.append(full_record.test_code)
        parts.append("\n")
        if full_record.entry_point:
            parts.append(f"\ncheck({full_record.entry_point})\n")
            
    combined_script = "".join(parts).strip()
    numbered_script = add_line_numbers(combined_script)
    
    ws.append([
        r.trace_id,
        r.dataset,
        r.problem_summary,
        numbered_script,
        r.traceback or r.stderr or "No traceback",
        "", 
        "", 
        ""  
    ])

for row in ws.iter_rows(min_row=2, max_row=len(eligible_records)+1):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 40
ws.column_dimensions['D'].width = 75
ws.column_dimensions['E'].width = 50
ws.column_dimensions['F'].width = 50
ws.column_dimensions['G'].width = 25
ws.column_dimensions['H'].width = 40

output_path = Path("annotation/Annotator_Task_100.xlsx")
wb.save(output_path)
print(f"Generated updated 100-case Excel sheet at {output_path}")
