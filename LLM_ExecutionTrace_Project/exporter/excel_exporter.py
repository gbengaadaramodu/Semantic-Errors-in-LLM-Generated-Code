"""
Exports execution results to an Excel workbook for manual annotation.

Creates a professionally formatted XLSX file with data validation
dropdowns on the Annotator columns, alternating row colours,
frozen header pane, and auto-width columns.
"""

import logging
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from models import ExecutionResult, ErrorCategory

logger = logging.getLogger(__name__)

# Column headers — must match the annotation schema
HEADERS = [
    "TraceID", "Dataset", "TaskID", "Problem Summary", "Execution Status",
    "stdout", "stderr", "Traceback", "Execution Time", "Parser Prediction",
    "Annotator1", "Annotator2", "Annotator3",
]


class AnnotationWorkbook:
    """Generates a formatted XLSX workbook for human annotation."""

    def __init__(self, output_path: Path):
        self.output_path = Path(output_path)

    def generate(self, results: List[ExecutionResult]) -> Path:
        """Create and save the annotation workbook."""
        logger.info("Generating Excel workbook → %s", self.output_path)

        wb = Workbook()
        ws = wb.active
        ws.title = "Validation_Corpus"

        # ── Header row ────────────────────────
        ws.append(HEADERS)

        # ── Sort: HE first, then MB, by trace_id ──
        sorted_results = sorted(
            results,
            key=lambda r: (0 if r.dataset == "HumanEval" else 1, r.trace_id),
        )

        # ── Data rows ─────────────────────────
        for r in sorted_results:
            ev = r.evidence
            ws.append([
                r.trace_id,
                r.dataset,
                r.task_id,
                r.problem_summary,
                r.execution_status,
                ev.stdout,
                ev.stderr,
                ev.traceback,
                f"{ev.execution_time:.3f}",
                r.parser_prediction,
                "",  # Annotator1
                "",  # Annotator2
                "",  # Annotator3
            ])

        # ── Header formatting ─────────────────
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(
            start_color="1F4E79", end_color="1F4E79", fill_type="solid",
        )
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # ── Freeze panes ──────────────────────
        ws.freeze_panes = "A2"

        # ── Auto-width (capped at 50) ─────────
        for col_cells in ws.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

        # ── Alternating row colours ───────────
        light_fill = PatternFill(
            start_color="F2F2F2", end_color="F2F2F2", fill_type="solid",
        )
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            if row_idx % 2 == 0:
                for cell in row:
                    cell.fill = light_fill

        # ── Data validation (Annotator dropdowns) ──
        category_values = ",".join(cat.value for cat in ErrorCategory)
        dv = DataValidation(
            type="list",
            formula1=f'"{category_values}"',
            allow_blank=True,
        )
        dv.error = "Please select a valid error category."
        dv.errorTitle = "Invalid Category"
        ws.add_data_validation(dv)
        last_row = len(sorted_results) + 1
        dv.add(f"K2:M{last_row}")

        # ── Save ──────────────────────────────
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.output_path)
        logger.info("Excel workbook saved (%d rows).", len(sorted_results))
        return self.output_path
