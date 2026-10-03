"""
Excel Exporter Module using `openpyxl`.

Generates the research-friendly `Objective2_Results.xlsx` workbook containing sheets:
    - Summary
    - Causal_Diagnosis
    - PDG_Evidence
    - Counterfactual_Results
    - Fault_Location
    - Runtime_Evidence
    - Ground_Truth
    - Error_Analysis
"""

import logging
from pathlib import Path
from typing import Any, List

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from objective2.config import Objective2Paths
from objective2.models import Objective2Result, Objective2Summary

logger = logging.getLogger(__name__)


class ExcelExporter:
    """
    Exports structured research results to a multi-sheet Excel workbook.
    """

    def __init__(self, paths: Objective2Paths) -> None:
        self.paths = paths

        # Styling definitions
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        self.center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def export(
        self, results: List[Objective2Result], summary: Objective2Summary
    ) -> Path:
        self.paths.ensure_directories()
        excel_path = self.paths.reports_dir / "Objective2_Results.xlsx"

        wb = openpyxl.Workbook()
        # Remove default sheet
        default_sheet = wb.active

        self._build_summary_sheet(wb, summary)
        self._build_causal_diagnosis_sheet(wb, results)
        self._build_pdg_evidence_sheet(wb, results)
        self._build_counterfactual_results_sheet(wb, results)
        self._build_fault_location_sheet(wb, results)
        self._build_runtime_evidence_sheet(wb, results)
        self._build_ground_truth_sheet(wb, results)
        self._build_error_analysis_sheet(wb, results)

        wb.remove(default_sheet)
        wb.save(excel_path)
        logger.info("Saved Excel report to %s", excel_path)
        return excel_path

    def _style_headers(self, ws: Any, headers: List[str]) -> None:
        ws.append(headers)
        ws.row_dimensions[1].height = 25
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def _auto_fit_columns(self, ws: Any, max_w: int = 45) -> None:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_w)

    def _build_summary_sheet(
        self, wb: openpyxl.Workbook, summary: Objective2Summary
    ) -> None:
        ws = wb.create_sheet("Summary")
        ws.append(["Objective 2: Causal Root Cause Analysis Summary Metrics"])

        title_cell = ws.cell(row=1, column=1)
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F4E78")

        rows = [
            ["Metric", "Value"],
            ["Total Corpus Size", summary.total_corpus_size],
            ["Eligible Cases", summary.eligible_cases],
            ["Analyzed Cases", summary.analyzed_cases],
            ["Skipped Cases", summary.skipped_cases],
            ["Unresolved Cases", summary.unresolved_cases],
            ["Analysis Failures", summary.failed_cases],
            ["Root Cause Accuracy", summary.root_cause_accuracy],
            ["Fault Location Accuracy", summary.fault_location_accuracy],
            ["Critical Variable Accuracy", summary.critical_variable_accuracy],
            ["Top-1 Fault Accuracy", summary.top_1_fault_accuracy],
            ["Top-3 Fault Accuracy", summary.top_3_fault_accuracy],
            ["Top-5 Fault Accuracy", summary.top_5_fault_accuracy],
            ["Counterfactual Causal Success Rate", summary.counterfactual_success_rate],
            ["Ground Truth Status", summary.ground_truth_status],
        ]

        for r in rows:
            ws.append(r)

        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 25

    def _build_causal_diagnosis_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("Causal_Diagnosis")
        headers = [
            "Trace ID",
            "Dataset",
            "Task ID",
            "Error Category",
            "Analysis Status",
            "Skip Reason",
            "Root Cause Status",
            "Candidate Explanation",
            "Causal Variable",
            "Fault Statement",
            "Fault Line",
            "Counterfactual Validated",
            "Confidence",
        ]
        self._style_headers(ws, headers)

        for r in results:
            ws.append([
                r.trace_id,
                r.dataset,
                r.task_id,
                r.error_category,
                r.analysis_status.value,
                r.skip_reason or "",
                r.root_cause.root_cause_status if r.root_cause else "",
                r.root_cause.candidate_explanation if r.root_cause else "",
                r.root_cause.causal_variable if r.root_cause else "",
                r.root_cause.fault_statement if r.root_cause else "",
                r.root_cause.fault_line if r.root_cause else "",
                r.root_cause.counterfactual_validated if r.root_cause else False,
                r.root_cause.confidence if r.root_cause else 0.0,
            ])

        self._auto_fit_columns(ws)

    def _build_pdg_evidence_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("PDG_Evidence")
        headers = [
            "Trace ID",
            "Total PDG Nodes",
            "Total PDG Edges",
            "Failure Node IDs",
            "Executed Nodes",
            "Failure Line",
        ]
        self._style_headers(ws, headers)

        for r in results:
            if r.pdg:
                exec_count = sum(1 for n in r.pdg.nodes.values() if n.executed)
                ws.append([
                    r.trace_id,
                    len(r.pdg.nodes),
                    len(r.pdg.edges),
                    str(r.pdg.failure_node_ids),
                    exec_count,
                    r.failure_line or "",
                ])

        self._auto_fit_columns(ws)

    def _build_counterfactual_results_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("Counterfactual_Results")
        headers = [
            "Trace ID",
            "Candidate Rank",
            "Target Variable",
            "Target Line",
            "Observed Value",
            "Counterfactual Value",
            "Original Outcome",
            "Counterfactual Outcome",
            "Causal Effect",
            "Execution Time (s)",
        ]
        self._style_headers(ws, headers)

        for r in results:
            for cf in r.counterfactual_results:
                ws.append([
                    r.trace_id,
                    cf.candidate_rank,
                    cf.intervention.target_variable,
                    cf.intervention.target_line,
                    str(cf.intervention.observed_value),
                    str(cf.intervention.counterfactual_value),
                    cf.original_outcome,
                    cf.counterfactual_outcome,
                    cf.causal_effect.value,
                    cf.execution_time,
                ])

        self._auto_fit_columns(ws)

    def _build_fault_location_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("Fault_Location")
        headers = [
            "Trace ID",
            "Dataset",
            "File",
            "Line Number",
            "Statement Code",
            "Function Name",
            "Causal Variable",
            "Evidence Score",
            "Confidence",
        ]
        self._style_headers(ws, headers)

        for r in results:
            if r.fault_location:
                fl = r.fault_location
                ws.append([
                    r.trace_id,
                    r.dataset,
                    fl.file,
                    fl.line_number,
                    fl.statement_code,
                    fl.function_name,
                    fl.causal_variable,
                    fl.evidence_score,
                    fl.confidence,
                ])

        self._auto_fit_columns(ws)

    def _build_runtime_evidence_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("Runtime_Evidence")
        headers = [
            "Trace ID",
            "Execution Status",
            "Exception Type",
            "Exception Message",
            "Runtime Events Count",
            "Stdout Snippet",
            "Stderr Snippet",
        ]
        self._style_headers(ws, headers)

        for r in results:
            rec = r.input_record
            if rec:
                ws.append([
                    r.trace_id,
                    rec.execution_status,
                    rec.exception_type,
                    rec.exception_message,
                    r.runtime_events_count,
                    rec.stdout[:100],
                    rec.stderr[:100],
                ])

        self._auto_fit_columns(ws)

    def _build_ground_truth_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("Ground_Truth")
        headers = [
            "Trace ID",
            "System Fault Line",
            "Ground Truth Fault Line",
            "Matched",
        ]
        self._style_headers(ws, headers)

        for r in results:
            fl_line = r.fault_location.line_number if r.fault_location else None
            ws.append([
                r.trace_id,
                fl_line or "",
                r.ground_truth_fault_line or "",
                r.ground_truth_matched if r.ground_truth_matched is not None else "PENDING",
            ])

        self._auto_fit_columns(ws)

    def _build_error_analysis_sheet(
        self, wb: openpyxl.Workbook, results: List[Objective2Result]
    ) -> None:
        ws = wb.create_sheet("Error_Analysis")
        headers = [
            "Trace ID",
            "Analysis Status",
            "Skip / Failure Rationale",
        ]
        self._style_headers(ws, headers)

        for r in results:
            if r.analysis_status != "ANALYZED":
                ws.append([
                    r.trace_id,
                    r.analysis_status.value,
                    r.skip_reason or "Unresolved execution path",
                ])

        self._auto_fit_columns(ws)
