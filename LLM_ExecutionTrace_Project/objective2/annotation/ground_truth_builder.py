"""
Ground Truth Builder Module.

Generates the research-friendly `objective2_ground_truth.xlsx` Excel template
and loads human annotation ground truth for verification.
"""

import logging
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)


class GroundTruthBuilder:
    """
    Creates and parses the objective2_ground_truth.xlsx ground-truth workbook.
    """

    HEADERS = [
        "Task ID",
        "Dataset",
        "Error Category",
        "Generated Code",
        "Human Root Cause",
        "Human Fault Line",
        "Human Fault Statement",
        "Human Causal Variable",
        "Evidence",
        "Annotator 1 ID",
        "Annotator 2 ID",
        "Annotator 3 ID",
        "Consensus Status",
    ]

    def create_template(self, excel_path: Path) -> Path:
        """Create empty ground truth Excel workbook with proper styling."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ground_Truth"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(self.HEADERS)
        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment

        # Set column widths
        widths = [12, 12, 20, 40, 35, 15, 30, 20, 30, 15, 15, 15, 18]
        for idx, w in enumerate(widths, 1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = w

        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(excel_path)
        logger.info("Created ground truth Excel template at %s", excel_path)
        return excel_path

    def load_ground_truth(self, excel_path: Path) -> Dict[str, dict]:
        """
        Load ground truth data indexed by Task ID.
        """
        if not excel_path.exists():
            return {}

        gt_dict: Dict[str, dict] = {}
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) <= 1:
                return {}

            headers = [str(h) for h in rows[0]]
            for r in rows[1:]:
                row_map = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                task_id = str(row_map.get("Task ID") or "").strip()
                if task_id and row_map.get("Human Fault Line") is not None:
                    gt_dict[task_id] = row_map
        except Exception as exc:
            logger.error("Failed loading ground truth Excel %s: %s", excel_path, exc)

        return gt_dict
