import json
import openpyxl
from openpyxl.styles import Alignment, Font
from pathlib import Path
from objective2.config import Objective2Config
from objective2.pipeline import Objective2Pipeline

config = Objective2Config()
pipe = Objective2Pipeline(config)
records = pipe.corpus_loader.load_records()
eligible_records = pipe._select_cases(records, error_type_filter=None, dataset_filter=None, task_id_filter=None, limit=None)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Annotation_Task"

headers = [
    "Trace ID", 
    "Dataset", 
    "Problem Summary (Prompt)", 
    "Generated Code (with Line Numbers)", 
    "Test Suite / Assertions",
    "Traceback / Error", 
    "-> YOUR ANNOTATION: Fault Line (Number)", 
    "-> YOUR ANNOTATION: Causal Variable", 
    "-> YOUR ANNOTATION: Bug Description"
]
ws.append(headers)

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = Font(bold=True)
    if "YOUR ANNOTATION" in header:
        cell.font = Font(bold=True, color="0000FF")

def add_line_numbers(code: str) -> str:
    if not code:
        return ""
    lines = code.splitlines()
    return "\n".join([f"{i+1:02d} | {line}" for i, line in enumerate(lines)])

for r in eligible_records:
    full_record = pipe.trace_loader.enrich_record(r)
    numbered_code = add_line_numbers(full_record.generated_code)
    
    test_suite = full_record.test_code or ""
    if full_record.test_imports:
        test_suite = full_record.test_imports + "\n\n" + test_suite
        
    ws.append([
        r.trace_id,
        r.dataset,
        r.problem_summary,
        numbered_code,
        test_suite,
        r.traceback or r.stderr or "No traceback",
        "", 
        "", 
        ""  
    ])

for row in ws.iter_rows(min_row=2, max_row=len(eligible_records)+1):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 40
ws.column_dimensions['D'].width = 60
ws.column_dimensions['E'].width = 50
ws.column_dimensions['F'].width = 50
ws.column_dimensions['G'].width = 20
ws.column_dimensions['H'].width = 25
ws.column_dimensions['I'].width = 40

output_path = Path("annotation/Annotator_Task_v2.xlsx")
wb.save(output_path)
print(f"Generated updated Excel sheet at {output_path}")
